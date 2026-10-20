"""Tests for src/extractors.py. Sample files are built in memory."""

from io import BytesIO

import docx
import pymupdf
import pytest

from src import extractors


def make_pdf(pages: list[str]) -> bytes:
    """Return a PDF with one page per string."""
    with pymupdf.open() as pdf:
        for text in pages:
            pdf.new_page().insert_text((72, 72), text)
        return pdf.tobytes()


def make_docx(paragraphs: list[str], table: list[list[str]]) -> bytes:
    """Return a DOCX with the given paragraphs followed by one table."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    word_table = document.add_table(rows=len(table), cols=len(table[0]))
    for row, values in zip(word_table.rows, table):
        for cell, value in zip(row.cells, values):
            cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_text_normalises_whitespace():
    raw = "  Hello\t\tworld  \r\n\r\n\r\n\r\nSecond   line  "
    assert extractors.extract_text(raw) == "Hello world\n\nSecond line"


@pytest.mark.parametrize("raw", ["", "   ", "\n\t\r\n"])
def test_extract_text_empty_input_returns_empty_string(raw):
    assert extractors.extract_text(raw) == ""


def test_extract_file_decodes_txt_with_bom():
    data = "﻿Meeting notes:  ship on Friday".encode("utf-8")
    assert extractors.extract_file("notes.TXT", data) == "Meeting notes: ship on Friday"


def test_extract_pdf_joins_pages_with_blank_line():
    data = make_pdf(["First page text", "Second page text"])
    assert extractors.extract_pdf(data) == "First page text\n\nSecond page text"


def test_extract_pdf_skips_blank_pages():
    data = make_pdf(["Only page with text", ""])
    assert extractors.extract_pdf(data) == "Only page with text"


def test_extract_pdf_rejects_damaged_file():
    with pytest.raises(ValueError, match="couldn't read this PDF"):
        extractors.extract_pdf(b"not a pdf")


def test_extract_docx_includes_paragraphs_and_table_cells():
    data = make_docx(["Project update", "All on track."], [["Task", "Owner"], ["Design", "Sam"]])
    assert extractors.extract_docx(data) == "Project update\nAll on track.\nTask | Owner\nDesign | Sam"


def test_extract_docx_rejects_damaged_file():
    with pytest.raises(ValueError, match="couldn't read this Word document"):
        extractors.extract_docx(b"not a docx")


@pytest.mark.parametrize("file_name", ["data.xlsx", "README", "photo.png", "scan.jpg"])
def test_extract_file_rejects_unsupported_type(file_name):
    with pytest.raises(ValueError, match="isn't a supported file type"):
        extractors.extract_file(file_name, b"anything")
