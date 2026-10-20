"""Turn each supported input type into plain text."""

import re
from io import BytesIO
from pathlib import Path

import docx
import pymupdf

from src import generator

TEXT_EXTENSIONS = {"txt", "md"}
IMAGE_MEDIA_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}


def extract_text(text: str) -> str:
    """Return text with tidy whitespace: single spaces, no trailing spaces, at most one blank line in a row."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract_pdf(data: bytes) -> str:
    """Return the text of every page in a PDF, pages separated by a blank line."""
    try:
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            pages = [extract_text(page.get_text()) for page in pdf]
    except (pymupdf.FileDataError, RuntimeError) as error:
        raise ValueError("We couldn't read this PDF. It may be damaged or password-protected.") from error
    return "\n\n".join(page for page in pages if page)


def extract_docx(data: bytes) -> str:
    """Return paragraph text followed by table rows (cells joined with ' | ') from a DOCX file."""
    try:
        document = docx.Document(BytesIO(data))
    except Exception as error:  # python-docx raises several unrelated types for bad files
        raise ValueError("We couldn't read this Word document. Please check it opens in Word.") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return extract_text("\n".join(parts))


def extract_image(data: bytes, media_type: str) -> str:
    """Return a text description of an image, produced by the vision model."""
    # No OCR: the vision model reads text and describes layout in one step.
    return extract_text(generator.describe_image(data, media_type))


def extract_file(file_name: str, data: bytes) -> str:
    """Pick the right extractor from the file extension and return its text."""
    extension = file_extension(file_name)
    if extension in TEXT_EXTENSIONS:
        return extract_text(data.decode("utf-8-sig", errors="replace"))
    if extension == "pdf":
        return extract_pdf(data)
    if extension == "docx":
        return extract_docx(data)
    if extension in IMAGE_MEDIA_TYPES:
        return extract_image(data, IMAGE_MEDIA_TYPES[extension])

    supported = ", ".join(sorted(TEXT_EXTENSIONS | IMAGE_MEDIA_TYPES.keys() | {"pdf", "docx"}))
    raise ValueError(f"Sorry, '{file_name}' isn't a supported file type. Please upload one of: {supported}.")


def file_extension(file_name: str) -> str:
    """Return the lower-case extension without the dot, e.g. 'pdf'."""
    return Path(file_name).suffix.lower().lstrip(".")

